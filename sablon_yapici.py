import os
from docx import Document

# Örneklerden birini baz alarak şablon oluşturalım
kaynak_ornek = ""
for root, dirs, files in os.walk("arsiv"):
    for f in files:
        if f.endswith(".docx") and not f.startswith("~$"):
            kaynak_ornek = os.path.join(root, f)
            break
    if kaynak_ornek:
        break

if not kaynak_ornek:
    print("HATA: Arşivde örnek docx bulunamadı!")
    exit(1)

# Boş şablonu orijinal stil özellikleriyle inşa ediyoruz
yeni_doc = Document(kaynak_ornek)

# İçeriği şablon etiketleriyle değiştiriyoruz
# Belgeyi temizle
for p in list(yeni_doc.paragraphs):
    p._element.getparent().remove(p._element)

# Temiz belgeye şablon yapısını ekle
yeni_doc.add_paragraph("{{ tarih }}")
yeni_doc.add_paragraph("{{ sayi }}")
yeni_doc.add_paragraph("Konu: {{ konu }}")
yeni_doc.add_paragraph("")
yeni_doc.add_paragraph("{{ hitap }}")
yeni_doc.add_paragraph("")
yeni_doc.add_paragraph("{{ kapsam }}")
yeni_doc.add_paragraph("")
yeni_doc.add_paragraph("{{ odeme_cumlesi }}")
yeni_doc.add_paragraph("")
yeni_doc.add_paragraph("{{ kapanis }}")
yeni_doc.add_paragraph("")

# İmza Bloğu (Dinamik)
p_imza = yeni_doc.add_paragraph("{% for imzaci in imzacilar %}{{ imzaci.ad_soyad }}\n{{ imzaci.unvan }}\n{% endfor %}")

yeni_doc.add_paragraph("")
yeni_doc.add_paragraph("{{ firma_unvan }}")
yeni_doc.add_paragraph("{{ banka }}")
yeni_doc.add_paragraph("IBAN: {{ iban }}")

os.makedirs("sablonlar", exist_ok=True)
yeni_doc.save("sablonlar/sablon.docx")
print("SABLON_BASARIYLA_OLUSTURULDU")